import os
import io
import time
import uuid
import asyncio
import pytest
from httpx import AsyncClient, ASGITransport
from PIL import Image
from pypdf import PdfWriter, PageObject
import docx
import openpyxl

from app.main import app
from app.core.config import settings


# --- Fixtures for Valid Sample Files ---

def create_sample_pdf(num_pages: int = 3) -> bytes:
    """Membuat file PDF valid multi-halaman in-memory."""
    writer = PdfWriter()
    for _ in range(num_pages):
        page = PageObject.create_blank_page(width=300, height=300)
        writer.add_page(page)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def create_sample_image(format: str = "PNG") -> bytes:
    """Membuat file gambar valid in-memory."""
    img = Image.new("RGB", (200, 200), color="blue")
    buf = io.BytesIO()
    img.save(buf, format=format)
    return buf.getvalue()


def create_sample_docx() -> bytes:
    """Membuat file Word DOCX valid in-memory."""
    doc = docx.Document()
    doc.add_heading("SmartDoc Test Document", 0)
    doc.add_paragraph("Ini adalah isi paragraf dokumen uji coba.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def wait_for_job_done(client: AsyncClient, job_id: str, max_wait: float = 15.0):
    """Helper untuk polling status background job hingga status 'done'."""
    t_start = time.time()
    while time.time() - t_start < max_wait:
        res = await client.get(f"/api/jobs/{job_id}")
        assert res.status_code == 200
        data = res.json()["data"]
        if data["status"] == "done":
            return data
        if data["status"] == "failed":
            raise RuntimeError(f"Job {job_id} gagal dengan error: {data.get('error')}")
        await asyncio.sleep(0.4)
    raise TimeoutError(f"Job {job_id} melebihi batas waktu tunggu {max_wait} detik")


# --- Test Cases End-to-End untuk Seluruh Fitur ---

@pytest.mark.asyncio
async def test_01_feature_compressor_image():
    """1. Test Fitur Compressor: Kompresi Gambar PNG"""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        png_bytes = create_sample_image("PNG")
        files = {"file": ("test_image.png", png_bytes, "image/png")}
        res = await client.post("/api/compressor/process", files=files)
        assert res.status_code == 202
        json_data = res.json()
        assert json_data["success"] is True
        job_id = json_data["data"]["job_id"]
        
        job_data = await wait_for_job_done(client, job_id)
        assert job_data["status"] == "done"
        assert job_data["progress"] == 100


@pytest.mark.asyncio
async def test_02_feature_compressor_pdf():
    """2. Test Fitur Compressor: Kompresi Dokumen PDF"""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        pdf_bytes = create_sample_pdf(2)
        files = {"file": ("test_doc.pdf", pdf_bytes, "application/pdf")}
        res = await client.post("/api/compressor/process", files=files)
        assert res.status_code == 202
        json_data = res.json()
        assert json_data["success"] is True
        job_id = json_data["data"]["job_id"]
        
        job_data = await wait_for_job_done(client, job_id)
        assert job_data["status"] == "done"


@pytest.mark.asyncio
async def test_03_feature_splitter_info_and_process():
    """3. Test Fitur Splitter: Info Halaman & Ekstrak Rentang Halaman"""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        pdf_bytes = create_sample_pdf(4)
        
        # A. Info Endpoint
        files = {"file": ("split_test.pdf", pdf_bytes, "application/pdf")}
        info_res = await client.post("/api/splitter/info", files=files)
        assert info_res.status_code == 200
        info_data = info_res.json()["data"]
        assert info_data["total_pages"] == 4
        
        # B. Process Split Endpoint
        split_files = {"file": ("split_test.pdf", pdf_bytes, "application/pdf")}
        split_data = {
            "split_mode": "extract_range",
            "range_expression": "1-2",
            "chunk_size": 1,
            "output_format": "pdf",
        }
        res = await client.post("/api/splitter/process", files=split_files, data=split_data)
        assert res.status_code == 202
        job_id = res.json()["data"]["job_id"]
        
        job_data = await wait_for_job_done(client, job_id)
        assert job_data["status"] == "done"


@pytest.mark.asyncio
async def test_04_feature_merger_info_and_process():
    """4. Test Fitur Merger: Penggabungan Beberapa PDF (Sequence Mode)"""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        pdf1 = create_sample_pdf(1)
        pdf2 = create_sample_pdf(2)
        
        # A. Merger Info Endpoint
        files_info = {"file": ("file1.pdf", pdf1, "application/pdf")}
        info_res = await client.post("/api/merger/info", files=files_info)
        assert info_res.status_code == 200
        assert info_res.json()["data"]["total_pages"] == 1
        
        # B. Merger Process Endpoint
        files = [
            ("files", ("part1.pdf", pdf1, "application/pdf")),
            ("files", ("part2.pdf", pdf2, "application/pdf")),
        ]
        data = {
            "merge_mode": "sequence",
            "insert_position": "end",
            "after_page": 1,
            "custom_title": "Merged_Document",
        }
        res = await client.post("/api/merger/process", files=files, data=data)
        assert res.status_code == 202
        job_id = res.json()["data"]["job_id"]
        
        job_data = await wait_for_job_done(client, job_id)
        assert job_data["status"] == "done"


@pytest.mark.asyncio
async def test_05_feature_watermark_and_security():
    """5. Test Fitur Watermark: Cap Air Teks & Enkripsi Password"""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        pdf_bytes = create_sample_pdf(2)
        files = {"file": ("watermark_target.pdf", pdf_bytes, "application/pdf")}
        data = {
            "watermark_text": "RAHASIA / CONFIDENTIAL",
            "opacity": 0.3,
            "angle": 45.0,
            "font_size": 36,
            "color_hex": "#ff0000",
            "password": "mypassword123",
        }
        res = await client.post("/api/watermark/process", files=files, data=data)
        assert res.status_code == 202
        job_id = res.json()["data"]["job_id"]
        
        job_data = await wait_for_job_done(client, job_id)
        assert job_data["status"] == "done"


@pytest.mark.asyncio
async def test_06_feature_converter_images_to_pdf():
    """6. Test Fitur Converter: Konversi Gambar PNG menjadi PDF"""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        png_bytes = create_sample_image("PNG")
        files = {"file": ("photo.png", png_bytes, "image/png")}
        data = {"output_format": "pdf"}
        res = await client.post("/api/converter/process", files=files, data=data)
        assert res.status_code == 202
        job_id = res.json()["data"]["job_id"]
        
        job_data = await wait_for_job_done(client, job_id)
        assert job_data["status"] == "done"


@pytest.mark.asyncio
async def test_07_feature_documents_history_and_batch():
    """7. Test Fitur Riwayat Dokumen: List, Single Download, Batch ZIP & Batch Delete"""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # A. List Documents
        res_list = await client.get("/api/documents")
        assert res_list.status_code == 200
        data_list = res_list.json()["data"]["items"]
        assert len(data_list) > 0
        
        target_doc = data_list[0]
        doc_id = target_doc["id"]
        
        # B. Single Download
        res_dl = await client.get(f"/api/documents/{doc_id}/download")
        assert res_dl.status_code in (200, 404)  # 200 if file generated, 404 if failed job
        
        # C. Batch Download ZIP (find docs with existing output files)
        valid_ids = [d["id"] for d in data_list if d.get("status") == "done"][:2]
        if valid_ids:
            res_batch = await client.post("/api/documents/batch-download", json={"ids": valid_ids})
            assert res_batch.status_code == 200
            assert res_batch.headers["content-type"] == "application/zip"
            assert len(res_batch.content) > 0
        
        # D. Batch Delete
        res_del = await client.post("/api/documents/batch-delete", json={"ids": [doc_id]})
        assert res_del.status_code == 200
        assert res_del.json()["success"] is True


@pytest.mark.asyncio
async def test_08_error_states_validation():
    """8. Test Penanganan State Error (Invalid input, bad range, corrupted files)"""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # A. Format tidak didukung (misal .exe)
        bad_file = {"file": ("malicious.exe", b"MZ\x90\x00", "application/octet-stream")}
        res_bad = await client.post("/api/converter/process", files=bad_file, data={"output_format": "pdf"})
        assert res_bad.status_code == 400
        assert res_bad.json()["success"] is False

        # B. Splitter dengan rentang halaman yang melampaui total halaman
        pdf_bytes = create_sample_pdf(2)
        files = {"file": ("short.pdf", pdf_bytes, "application/pdf")}
        split_data = {
            "split_mode": "extract_range",
            "range_expression": "10-20",  # PDF hanya punya 2 halaman
            "chunk_size": 1,
            "output_format": "pdf",
        }
        res_range = await client.post("/api/splitter/process", files=files, data=split_data)
        assert res_range.status_code == 202
        job_id = res_range.json()["data"]["job_id"]
        
        # Polling job status dan verifikasi transisi ke state 'failed' dengan pesan yang benar
        await asyncio.sleep(1.0)
        job_res = await client.get(f"/api/jobs/{job_id}")
        assert job_res.status_code == 200
        job_data = job_res.json()["data"]
        assert job_data["status"] == "failed"
        assert "tidak valid" in job_data["error"].lower()
